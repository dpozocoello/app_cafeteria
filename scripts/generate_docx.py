import os
import re
import zlib
import base64
import requests
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def kroki_encode(data: str) -> str:
    """Comprime el texto con zlib y lo codifica en base64 url-safe para la API de Kroki."""
    compressed = zlib.compress(data.encode('utf-8'), level=9)
    encoded = base64.urlsafe_b64encode(compressed).decode('utf-8')
    return encoded

def render_mermaid_diagram(mermaid_code: str) -> str:
    """Descarga el diagrama Mermaid renderizado como PNG usando Kroki."""
    print("Renderizando diagrama Mermaid con Kroki...")
    try:
        encoded = kroki_encode(mermaid_code)
        kroki_url = f"https://kroki.io/mermaid/png/{encoded}"
        response = requests.get(kroki_url, timeout=20)
        
        if response.status_code == 200:
            os.makedirs("scripts/temp_images", exist_ok=True)
            # Generar un nombre único basado en hash
            filename = f"scripts/temp_images/mermaid_{hash(mermaid_code)}.png"
            with open(filename, "wb") as f:
                f.write(response.content)
            return filename
        else:
            print(f"Error de Kroki ({response.status_code}): {response.text}")
    except Exception as e:
        print(f"Excepción al renderizar diagrama: {e}")
    return None

def set_cell_shading(cell, color_hex):
    """Aplica color de fondo (shading) a una celda de tabla en Word."""
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table):
    """Aplica bordes modernos a una tabla (sin bordes verticales, horizontales finos)."""
    tblPr = table._tbl.tblPr
    borders_xml = (
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="94A3B8"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(parse_xml(borders_xml))

def add_markdown_runs(p, text):
    """Parsea el texto en busca de markdown en línea (**negrita**, *cursiva*, `código`) y lo añade como runs."""
    # Regex para detectar formatos en línea manteniendo los delimitadores
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\[.*?\]\(.*?\))')
    parts = pattern.split(text)
    
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            content = part[2:-2]
            run = p.add_run(content)
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            content = part[1:-1]
            run = p.add_run(content)
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            content = part[1:-1]
            run = p.add_run(content)
            run.font.name = 'Courier New'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(199, 37, 78) # Color burdeos (#c7254e) para código
        elif part.startswith('[') and ']' in part and '(' in part and part.endswith(')'):
            match = re.match(r'\[(.*?)\]\((.*?)\)', part)
            if match:
                link_text = match.group(1)
                run = p.add_run(link_text)
                run.font.color.rgb = RGBColor(37, 99, 235) # Azul enlace (#2563eb)
                run.underline = True
            else:
                p.add_run(part)
        else:
            # Reemplazar fórmulas matemáticas simples para que se vean limpias
            clean_text = part.replace('$$\\text{Suma} = \\sum_{i=1}^{N} d_i \\times f_i$$', 'Suma = Σ (d_i × f_i)')
            clean_text = clean_text.replace('$$\\text{Residuo} = \\text{Suma} \\pmod{11}$$', 'Residuo = Suma Modulo 11')
            clean_text = clean_text.replace('$$\\text{Verificador} = 11 - \\text{Residuo}$$', 'Verificador = 11 - Residuo')
            p.add_run(clean_text)

def main():
    md_file_path = "DOCUMENTACION_TECNICA.md"
    docx_file_path = "DOCUMENTACION_TECNICA.docx"
    
    if not os.path.exists(md_file_path):
        print(f"Error: No se encontró {md_file_path}")
        return

    doc = Document()
    
    # ─── Configurar Estilos Predeterminados ───────────────────────────────────────
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(51, 65, 85) # Slate-700 (#334155)

    # ─── Generar Portada ─────────────────────────────────────────────────────────
    print("Creando portada...")
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(140)
    title_p.paragraph_format.space_after = Pt(12)
    title_run = title_p.add_run("COFFEEAPP v2.0")
    title_run.font.size = Pt(32)
    title_run.bold = True
    title_run.font.color.rgb = RGBColor(15, 23, 42) # Slate-900 (#0F172A)

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_p.paragraph_format.space_after = Pt(60)
    sub_run = subtitle_p.add_run("Documentación Técnica de Software")
    sub_run.font.size = Pt(16)
    sub_run.font.color.rgb = RGBColor(100, 116, 139) # Slate-500

    desc_p = doc.add_paragraph()
    desc_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc_p.paragraph_format.space_before = Pt(80)
    desc_run = desc_p.add_run(
        "Sistema de Gestión Comercial, Control de Kárdex por Recetas,\n"
        "Facturación Electrónica (SRI Ecuador), Seguridad ISO 27001\n"
        "y Cumplimiento de Privacidad LOPDP/GDPR\n\n"
        "Diseñado para Desarrolladores y Administradores de Sistemas"
    )
    desc_run.font.size = Pt(11)
    desc_run.font.italic = True
    desc_run.font.color.rgb = RGBColor(71, 85, 105) # Slate-600

    doc.add_page_break()

    # ─── Leer y Parsear DOCUMENTACION_TECNICA.md ────────────────────────────────
    print("Parseando DOCUMENTACION_TECNICA.md...")
    with open(md_file_path, "r", encoding="utf-8") as f:
        lines = f.readlines()
        
    i = 0
    n = len(lines)
    
    while i < n:
        line = lines[i]
        stripped = line.strip()
        
        # Saltar líneas vacías
        if not line.replace('\n', ''):
            i += 1
            continue
            
        # Bloques de Código (Mermaid o normal)
        if stripped.startswith('```'):
            lang = stripped[3:].strip()
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1 # omitir final de ```
            
            code_content = "".join(code_lines)
            
            if lang == 'mermaid':
                # Renderizar Mermaid e insertarlo como imagen
                img_path = render_mermaid_diagram(code_content)
                if img_path and os.path.exists(img_path):
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(8)
                    p.paragraph_format.space_after = Pt(8)
                    run = p.add_run()
                    run.add_picture(img_path, width=Inches(5.8))
                    
                    # Espaciado extra
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_after = Pt(12)
                    cap_run = p_cap.add_run("▲ Representación gráfica del diagrama renderizado")
                    cap_run.font.size = Pt(8.5)
                    cap_run.font.italic = True
                    cap_run.font.color.rgb = RGBColor(148, 163, 184) # Slate-400
                    
                    os.remove(img_path)
                else:
                    # Fallback si Kroki falla: insertar como texto de código
                    table = doc.add_table(rows=1, cols=1)
                    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cell = table.cell(0, 0)
                    set_cell_shading(cell, "F8FAFC")
                    p = cell.paragraphs[0]
                    run = p.add_run(code_content.rstrip())
                    run.font.name = 'Courier New'
                    run.font.size = Pt(8.5)
            else:
                # Código estándar (Python / Bash) - Caja gris con bordes
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cell = table.cell(0, 0)
                set_cell_shading(cell, "F1F5F9")
                
                tcPr = cell._tc.get_or_add_tcPr()
                tcBorders = parse_xml(
                    f'<w:tcBorders {nsdecls("w")}>'
                    f'  <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                    f'  <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                    f'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                    f'  <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
                    f'</w:tcBorders>'
                )
                tcPr.append(tcBorders)
                
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.left_indent = Pt(6)
                
                run = p.add_run(code_content.rstrip())
                run.font.name = 'Courier New'
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(51, 65, 85)
            continue
            
        # Encabezados
        if stripped.startswith('# '):
            h_text = stripped[2:].strip()
            p = doc.add_heading(level=1)
            p.paragraph_format.space_before = Pt(20)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.color.rgb = RGBColor(15, 23, 42)
            run.bold = True
            i += 1
            continue
        elif stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = doc.add_heading(level=2)
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.color.rgb = RGBColor(15, 23, 42)
            run.bold = True
            i += 1
            continue
        elif stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = doc.add_heading(level=3)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.font.color.rgb = RGBColor(71, 85, 105)
            run.bold = True
            i += 1
            continue
        elif stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(h_text)
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(100, 116, 139)
            i += 1
            continue

        # Tablas
        if stripped.startswith('|'):
            table_lines = []
            while i < n and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
                
            if len(table_lines) >= 2:
                headers = [c.strip() for c in table_lines[0].split('|')[1:-1]]
                row_lines = table_lines[2:] # omitir cabecera y separador
                
                table = doc.add_table(rows=len(row_lines) + 1, cols=len(headers))
                table.alignment = WD_ALIGN_PARAGRAPH.CENTER
                set_table_borders(table)
                
                # Cabecera
                hdr_cells = table.rows[0].cells
                for col_idx, text in enumerate(headers):
                    hdr_cells[col_idx].text = ""
                    p = hdr_cells[col_idx].paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run(text)
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.font.size = Pt(9.5)
                    set_cell_shading(hdr_cells[col_idx], "0F172A") # Navy Blue
                    
                # Filas
                for row_idx, r_line in enumerate(row_lines):
                    row_cells = table.rows[row_idx + 1].cells
                    columns = [c.strip() for c in r_line.split('|')[1:-1]]
                    for col_idx, text in enumerate(columns):
                        if col_idx < len(row_cells):
                            row_cells[col_idx].text = ""
                            p = row_cells[col_idx].paragraphs[0]
                            p.paragraph_format.space_before = Pt(3)
                            p.paragraph_format.space_after = Pt(3)
                            add_markdown_runs(p, text)
                            
                            # Cebra (filas alternas gris claro)
                            if row_idx % 2 == 1:
                                set_cell_shading(row_cells[col_idx], "F8FAFC")
            continue

        # Listas desordenadas
        if stripped.startswith('- ') or stripped.startswith('* ') or stripped.startswith('  - ') or stripped.startswith('  * '):
            content = stripped.lstrip('-* ').strip()
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2.5)
            if stripped.startswith('  '):
                p.paragraph_format.left_indent = Inches(0.5)
            add_markdown_runs(p, content)
            i += 1
            continue
            
        # Listas ordenadas
        if re.match(r'^\d+\.\s', stripped):
            match = re.match(r'^\d+\.\s(.*)', stripped)
            content = match.group(1).strip()
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2.5)
            add_markdown_runs(p, content)
            i += 1
            continue

        # Párrafo estándar
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        add_markdown_runs(p, stripped)
        i += 1

    doc.save(docx_file_path)
    print(f"Documento de Word generado exitosamente en: {docx_file_path}")

    # Limpieza de imágenes temporales
    if os.path.exists("scripts/temp_images"):
        for f in os.listdir("scripts/temp_images"):
            try:
                os.remove(os.path.join("scripts/temp_images", f))
            except:
                pass
        try:
            os.rmdir("scripts/temp_images")
        except:
            pass

if __name__ == "__main__":
    main()
